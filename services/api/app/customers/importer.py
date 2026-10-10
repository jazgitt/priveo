"""One-time load of the existing customers from the owner's spreadsheet.

It reads two sheets, "Customer Master" (one row for each vehicle) and
"Deep Spa Addons". A person in the second sheet who is not in the first has
taken Deep Spa as a standalone service (ADD-14): they are loaded as a
customer with a vehicle and no subscription. It can run again: a customer is matched by email and a
row by its serial number, and what is already loaded is left as it is.

The report names rows by sheet and serial number only. It never contains a
name, phone number, email or registration number.
"""
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.customers.models import WEEKDAYS, AddonRequest, Customer, Subscription, Vehicle
from app.directory.models import Society
from app.identity.models import Login
from app.platform.clock import today_ist

PLACEHOLDER_SOCIETY = "To be confirmed"
MASTER_SHEET = "Customer Master"
ADDON_SHEET = "Deep Spa Addons"

_EMAIL = re.compile(r"[^@\s]+@[^@\s]+\.[a-z]{2,}")
# DL 13 CM 0019, and the Bharat series 22 BH 1234 AA.
_REGISTRATION = re.compile(r"[A-Z]{2}\d{1,2}[A-Z]{0,3}\d{4}|\d{2}BH\d{4}[A-Z]{1,2}")
_PARKING = re.compile(r"(B\d+)\s*-\s*(\w+)", re.IGNORECASE)
_MONTHS = {m: i + 1 for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
_PAYMENT = {"paid": "Received", "pending": "Due"}


@dataclass
class ImportReport:
    applied: bool = False
    customers_created: int = 0
    customers_kept: int = 0
    vehicles_created: int = 0
    vehicles_kept: int = 0
    # Vehicles of customers who took only a standalone Deep Spa.
    standalone_vehicles_created: int = 0
    standalone_vehicles_kept: int = 0
    addons_created: int = 0
    addons_kept: int = 0
    # Loaded, but something needs the admin's attention.
    warnings: list[str] = field(default_factory=list)
    # Not loaded.
    errors: list[str] = field(default_factory=list)


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _rows(sheet) -> list[dict[str, object]]:
    """Each data row as a dict keyed by its lower-case column heading."""
    rows = sheet.iter_rows(values_only=True)
    headings = [_text(h).lower() for h in next(rows, ())]
    out = []
    for values in rows:
        if any(v not in (None, "") for v in values):
            out.append({h: v for h, v in zip(headings, values, strict=False) if h})
    return out


def split_name(name: str) -> tuple[str, str]:
    first, _, last = name.strip().rpartition(" ")
    return (first, last) if first else (last, "")


def mobile_of(value) -> str | None:
    first = _text(value).split("/")[0].strip()
    # A masked file hides digits with "x", as in 9xxxxx1234. It is kept as
    # given, so the load works on data that is safe to share.
    if re.fullmatch(r"[0-9xX]{10}", first) and "x" in first.lower():
        return first.lower()
    digits = re.sub(r"\D", "", first)
    digits = digits[-10:]
    return digits if len(digits) == 10 else None


def registration_of(value) -> str | None:
    compact = re.sub(r"[^A-Z0-9]", "", _text(value).upper())
    return compact if _REGISTRATION.fullmatch(compact) else None


def parking_of(value) -> tuple[str, str, bool]:
    """Basement, number, and whether the whole text was understood."""
    text = " ".join(_text(value).split())
    whole = _PARKING.fullmatch(text)
    if whole:
        return whole.group(1).upper(), whole.group(2), True
    start = re.match(r"B\d+\b", text, re.IGNORECASE)
    return (start.group(0).upper() if start else ""), "", False


def time_of(value) -> time | None:
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::\d{2})?", _text(value))
    return time(int(match.group(1)), int(match.group(2))) if match else None


def date_of(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


def addon_date_of(value, year: int) -> date | None:
    """Reads "4th Oct" or "7th Oct 5:30 PM". The sheet gives no year."""
    match = re.match(r"\s*(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]{3})", _text(value))
    if not match or match.group(2).lower() not in _MONTHS:
        return None
    try:
        return date(year, _MONTHS[match.group(2).lower()], int(match.group(1)))
    except ValueError:
        return None


def _compact(value) -> str:
    return re.sub(r"[^A-Z0-9]", "", _text(value).upper())


class _Importer:
    def __init__(self, db: Session, report: ImportReport):
        self.db = db
        self.report = report
        self.society: Society | None = None
        # For each email: the customer and the vehicles from the sheet, each
        # with the car number and model text the sheet gave.
        self.vehicles: dict[str, list[tuple[Vehicle, str, str]]] = {}
        self.customers: dict[str, Customer] = {}
        self.master_emails: set[str] = set()
        self.latest_start: date | None = None

    def warn(self, sheet: str, ref: str, message: str) -> None:
        self.report.warnings.append(f"{sheet} row {ref}: {message}")

    def error(self, sheet: str, ref: str, message: str) -> None:
        self.report.errors.append(f"{sheet} row {ref}: {message}")

    def _society(self) -> Society:
        if self.society is None:
            self.society = self.db.scalar(
                select(Society).where(Society.name == PLACEHOLDER_SOCIETY))
            if self.society is None:
                self.society = Society(name=PLACEHOLDER_SOCIETY)
                self.db.add(self.society)
                self.db.flush()
        return self.society

    def _customer(self, row: dict, sheet: str, ref: str, email: str) -> Customer | None:
        if email in self.customers:
            return self.customers[email]
        customer = self.db.scalar(select(Customer).where(Customer.email == email))
        if customer is not None:
            self.report.customers_kept += 1
        else:
            name = _text(row.get("customer name"))
            mobile = mobile_of(row.get("phone number"))
            if not name or mobile is None:
                self.error(sheet, ref, "the name or the 10-digit phone number is missing")
                return None
            if self.db.scalar(select(Login).where(Login.login_id == email)) is not None:
                self.error(sheet, ref, "the email already belongs to another login")
                return None
            first, last = split_name(name)
            if not last:
                self.warn(sheet, ref, "the name is one word, so the last name is empty")
            # No password yet: the customer sets one at first sign-in (AUTH-03).
            login = Login(role="customer", login_id=email, display_name=name)
            self.db.add(login)
            self.db.flush()
            customer = Customer(
                login_id=login.id, first_name=first, last_name=last, email=email,
                mobile=mobile, society_id=self._society().id,
                flat=_text(row.get("flat no.")), source="migration")
            self.db.add(customer)
            self.db.flush()
            self.report.customers_created += 1
        self.customers[email] = customer
        return customer

    def _new_vehicle(self, customer: Customer, row: dict, sheet: str, ref: str,
                     source_ref: str) -> Vehicle:
        car_number = _text(row.get("car number"))
        notes = []
        registration = registration_of(car_number)
        if registration is None:
            self.warn(sheet, ref, "the car number is missing or incomplete"
                      if not car_number or len(_compact(car_number)) < 8
                      else "the car number is not in a known format")
            if car_number:
                notes.append(f"Car number as given: {car_number}")
        elif self.db.scalar(select(Vehicle).where(Vehicle.registration == registration)):
            self.warn(sheet, ref, "the car number is already on another vehicle")
            notes.append(f"Car number as given: {car_number}")
            registration = None
        basement, number, understood = parking_of(row.get("parking number"))
        if not understood:
            self.warn(sheet, ref, "the parking is free text, kept as a note")
            notes.append(f"Parking as given: {' '.join(_text(row.get('parking number')).split())}")

        vehicle = Vehicle(
            customer_id=customer.id, registration=registration,
            model=_text(row.get("car model")), parking_basement=basement,
            parking_number=number, notes="\n".join(notes), source_ref=source_ref)
        self.db.add(vehicle)
        self.db.flush()
        return vehicle

    def master_row(self, row: dict) -> None:
        ref = _text(row.get("sl. no.")) or "?"
        email = _text(row.get("email")).lower()
        if not _EMAIL.fullmatch(email):
            return self.error(MASTER_SHEET, ref, "the email is missing or not valid")
        start = date_of(row.get("start date"))
        price = row.get("subscription price (inr)")
        model = _text(row.get("car model"))
        if start is None or not isinstance(price, int | float) or not model:
            return self.error(MASTER_SHEET, ref,
                              "the start date, the price or the car model is missing")
        customer = self._customer(row, MASTER_SHEET, ref, email)
        if customer is None:
            return None
        self.master_emails.add(email)
        self.latest_start = max(self.latest_start or start, start)

        source_ref = f"master:{ref}"
        car_number = _text(row.get("car number"))
        vehicle = self.db.scalar(select(Vehicle).where(Vehicle.source_ref == source_ref))
        if vehicle is not None:
            self.report.vehicles_kept += 1
            self.vehicles.setdefault(email, []).append((vehicle, car_number, model))
            return None

        vehicle = self._new_vehicle(customer, row, MASTER_SHEET, ref, source_ref)

        cutoff = time_of(row.get("customer cutoff time"))
        if cutoff is None:
            self.warn(MASTER_SHEET, ref, "no cutoff time")
        day = _text(row.get("interior preferred day")).capitalize() or None
        if day is not None and day not in WEEKDAYS:
            self.warn(MASTER_SHEET, ref, "the interior day is not a weekday, left empty")
            day = None
        self.db.add(Subscription(
            vehicle_id=vehicle.id, monthly_price_rupees=int(price), status="Active",
            first_service_date=start, cutoff_time=cutoff, interior_day=day))
        self.report.vehicles_created += 1
        self.vehicles.setdefault(email, []).append((vehicle, car_number, model))
        return None

    def _match_vehicle(self, email: str, car_number: str, model: str,
                       *, guess: bool = True) -> Vehicle | None:
        """The customer's vehicle with this car number or model. With guess,
        a customer's only vehicle is taken when neither matches."""
        mine = self.vehicles.get(email, [])
        wanted = _compact(car_number)
        if wanted:
            for vehicle, number, _ in mine:
                have = _compact(number)
                if have and (have == wanted or have.endswith(wanted) or wanted.endswith(have)):
                    return vehicle
        for vehicle, _, their_model in mine:
            if model and their_model.casefold() == model.casefold():
                return vehicle
        return mine[0][0] if guess and len(mine) == 1 else None

    def _standalone_vehicle(self, customer: Customer, row: dict, ref: str, email: str,
                            car_number: str, model: str) -> None:
        """The car of a customer who has no subscription (ADM-ADD-11)."""
        source_ref = f"deepspa-vehicle:{ref}"
        vehicle = self.db.scalar(select(Vehicle).where(Vehicle.source_ref == source_ref))
        if vehicle is not None:
            self.report.standalone_vehicles_kept += 1
        elif self._match_vehicle(email, car_number, model, guess=False) is not None:
            return None  # The same car as an earlier row of this customer.
        elif not model:
            return self.warn(ADDON_SHEET, ref, "no car model, so no vehicle was made")
        else:
            vehicle = self._new_vehicle(customer, row, ADDON_SHEET, ref, source_ref)
            self.report.standalone_vehicles_created += 1
        self.vehicles.setdefault(email, []).append((vehicle, car_number, model))
        return None

    def addon_row(self, row: dict, year: int) -> None:
        ref = _text(row.get("sl. no.")) or "?"
        email = _text(row.get("email")).lower()
        if not _EMAIL.fullmatch(email):
            return self.error(ADDON_SHEET, ref, "the email is missing or not valid")
        customer = self._customer(row, ADDON_SHEET, ref, email)
        if customer is None:
            return None
        car_number = _text(row.get("car number"))
        model = _text(row.get("car model"))
        if email not in self.master_emails:
            self._standalone_vehicle(customer, row, ref, email, car_number, model)
        source_ref = f"deepspa:{ref}"
        if self.db.scalar(select(AddonRequest).where(AddonRequest.source_ref == source_ref)):
            self.report.addons_kept += 1
            return None

        package = _text(row.get("package"))
        price = re.search(r"\((\d+)", package)
        service = package.split("(")[0].strip() or "Deep Spa"
        when_text = _text(row.get("date"))
        when = addon_date_of(when_text, year)
        if when is None:
            self.warn(ADDON_SHEET, ref, "the date could not be read, kept as a note")
        status_text = _text(row.get("status")).lower()
        if status_text not in _PAYMENT:
            self.warn(ADDON_SHEET, ref, "no payment status")
        vehicle = self._match_vehicle(email, car_number, model)
        if vehicle is None:
            self.warn(ADDON_SHEET, ref, "could not tell which of the customer's vehicles")

        notes = [n for n in (f"Date as given: {when_text}" if when_text else "",
                             _text(row.get("notes"))) if n]
        self.db.add(AddonRequest(
            customer_id=customer.id, vehicle_id=vehicle.id if vehicle else None,
            service=service, price_rupees=int(price.group(1)) if price else None,
            service_date=when, payment_status=_PAYMENT.get(status_text),
            notes="\n".join(notes), source_ref=source_ref))
        self.report.addons_created += 1
        return None


def import_workbook(db: Session, path: Path | str, *, apply: bool) -> ImportReport:
    """Loads the workbook. With apply=False nothing is saved: it is a trial
    run that reports what would happen."""
    workbook = load_workbook(path, data_only=True)
    report = ImportReport(applied=apply)
    run = _Importer(db, report)
    try:
        if MASTER_SHEET not in workbook.sheetnames:
            report.errors.append(f'There is no sheet named "{MASTER_SHEET}".')
        else:
            for row in _rows(workbook[MASTER_SHEET]):
                run.master_row(row)
        if ADDON_SHEET in workbook.sheetnames:
            year = (run.latest_start or today_ist()).year
            for row in _rows(workbook[ADDON_SHEET]):
                run.addon_row(row, year)
        if apply:
            db.commit()
        else:
            db.rollback()
    except Exception:
        db.rollback()
        raise
    return report
